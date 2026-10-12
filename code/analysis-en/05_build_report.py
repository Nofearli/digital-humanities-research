"""Generate a Chinese academic-style Word report from completed analysis results."""

from __future__ import annotations

import importlib.metadata
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from pipeline_common import (
    DATA_DIR, FIGURES_DIR, RAW_COMMENTS, RAW_VIDEOS, REPORTS_DIR, RESULTS_DIR,
    ensure_directories, read_json, sha256, utc_now, write_json,
)


TOPIC_CN = {
    0: "观影评价、价值判断与成长叙事",
    1: "动画品质、影评互动与观影推荐",
    2: "中国电影、好莱坞与产业竞争",
    3: "创作者互动与健康关怀",
    4: "剧情片段、角色与视觉解读",
    5: "“丑小鸭”隐喻、身份认同与成长",
    6: "东亚文化、发音与身份话语",
}


def set_run_font(run, size: float = 10.5, bold: bool = False) -> None:
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(size)
    run.bold = bold


def set_cell_text(cell, text: object) -> None:
    cell.text = str(text)
    for paragraph in cell.paragraphs:
        paragraph.paragraph_format.space_after = Pt(0)
        for run in paragraph.runs:
            set_run_font(run, 8.5)


def add_caption(document: Document, caption: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(6)
    run = paragraph.add_run(caption)
    set_run_font(run, 10)


def add_figure(document: Document, filename: str, caption: str, explanation: str = "", width: float = 15.5) -> None:
    path = FIGURES_DIR / filename
    if path.exists():
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.add_run().add_picture(str(path), width=Cm(width))
        add_caption(document, caption)
        if explanation:
            add_body(document, explanation)


def add_table(document: Document, frame: pd.DataFrame, columns: list[str], limit: int | None = None) -> None:
    display = frame.loc[:, columns].head(limit) if limit else frame.loc[:, columns]
    table = document.add_table(rows=1, cols=len(columns))
    table.style = "Table Grid"
    for cell, column in zip(table.rows[0].cells, columns):
        set_cell_text(cell, column)
    for _, row in display.iterrows():
        cells = table.add_row().cells
        for cell, column in zip(cells, columns):
            set_cell_text(cell, row[column])


def add_body(document: Document, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.first_line_indent = Cm(0.74)
    paragraph.paragraph_format.line_spacing = 1.45
    paragraph.paragraph_format.space_after = Pt(5)
    run = paragraph.add_run(text)
    set_run_font(run, 10.5)


def configure_styles(document: Document) -> None:
    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)
    for name in ("Heading 1", "Heading 2"):
        style = document.styles[name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    for section in document.sections:
        section.top_margin = Cm(2.54)
        section.bottom_margin = Cm(2.54)
        section.left_margin = Cm(3.0)
        section.right_margin = Cm(3.0)


def theme_sentence(topic_id: int, count: int, share: float) -> str:
    interpretation = {
        0: "这一最大主题把作品体验与个人成长、价值判断连在一起，说明评论者不只讨论技术层面的“好看”，也把影片置入成长、努力与自我理解的叙事框架中。",
        1: "该主题围绕动画效果、影评视频与推荐行为展开，呈现出作品评价与平台内容消费相互嵌套的特征：观众既评价影片，也回应讲解者或推荐者。",
        2: "该主题将作品置于中国电影、好莱坞和市场竞争的比较中，反映出跨文化受众会借由产业位置、制作标准和文化代表性来理解影片。",
        3: "该主题包含对视频创作者的问候、互动和健康关怀，属于 YouTube 评论区特有的关系性话语；它提醒研究者将“对电影的接受”与“对视频创作者的互动”区分开来。",
        4: "该主题聚焦具体剧情、角色和视觉片段，体现出受众以叙事情节与角色形象组织观影记忆，并以局部场景进入作品意义。",
        5: "该主题以“丑小鸭”等成长隐喻连接身份、差异和被认可的经验，显示观众会将动画叙事转化为可共鸣的个人处境。",
        6: "该主题涉及东亚文化、语言发音与身份称谓，表明跨语言传播中的文化识读、命名方式和区域身份同样构成讨论焦点。",
    }
    return f"主题 {topic_id}“{TOPIC_CN.get(topic_id, '其他主题')}”包含 {count:,} 条评论，占有效主题评论的 {share:.1%}。{interpretation.get(topic_id, '')}"


def main() -> None:
    ensure_directories()
    prep = read_json(DATA_DIR / "preparation_summary.json")
    topics = pd.read_csv(RESULTS_DIR / "topic_summary.csv", encoding="utf-8-sig")
    keywords = pd.read_csv(RESULTS_DIR / "topic_keywords.csv", encoding="utf-8-sig")
    quotes = pd.read_csv(RESULTS_DIR / "topic_representative_quotes.csv", encoding="utf-8-sig")
    sentiment = pd.read_csv(RESULTS_DIR / "sentiment_summary.csv", encoding="utf-8-sig")
    sentiment_keywords = pd.read_csv(RESULTS_DIR / "sentiment_keywords.csv", encoding="utf-8-sig")
    network = pd.read_csv(RESULTS_DIR / "network_nodes.csv", encoding="utf-8-sig")
    topic_meta = read_json(RESULTS_DIR / "topic_metadata.json")
    sentiment_meta = read_json(RESULTS_DIR / "sentiment_metadata.json")
    network_meta = read_json(RESULTS_DIR / "network_metadata.json")
    crosstab_path = RESULTS_DIR / "topic_sentiment_crosstab.csv"
    crosstab = pd.read_csv(crosstab_path, encoding="utf-8-sig") if crosstab_path.exists() else pd.DataFrame()

    english_total = prep["funnel"]["english_analysis_corpus"]
    visible_topics = topics[~topics["is_outlier"].astype(str).str.lower().eq("true")].sort_values("comment_count", ascending=False).copy()
    visible_topics["中文主题"] = visible_topics["topic_id"].map(TOPIC_CN).fillna(visible_topics["interpreted_label"])
    visible_total = int(visible_topics["comment_count"].sum())
    sentiment_map = dict(zip(sentiment["sentiment"], sentiment["share"]))
    sentiment_count = dict(zip(sentiment["sentiment"], sentiment["comment_count"]))
    positive_terms = "、".join(sentiment_keywords[sentiment_keywords["sentiment"].eq("positive")].head(8)["term"].tolist())
    negative_terms = "、".join(sentiment_keywords[sentiment_keywords["sentiment"].eq("negative")].head(8)["term"].tolist())
    top_network_terms = "、".join(network.head(8)["term"].tolist())

    document = Document()
    configure_styles(document)
    heading = document.add_heading("《哪吒之魔童闹海》英语 YouTube 评论的计算文本分析", 0)
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in heading.runs:
        set_run_font(run, 16, bold=True)
    subheading = document.add_paragraph("——主题、情感与语义网络的探索性研究报告")
    subheading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in subheading.runs:
        set_run_font(run, 11)

    document.add_heading("摘要", level=1)
    add_body(document, f"本研究以 YouTube 平台上《哪吒之魔童闹海》（Ne Zha 2）相关视频的公开顶层评论为材料，参照计算传播研究中“内容—情感—结构”的分析路径，考察英语受众如何讨论和评价这部动画电影。研究从 {prep['funnel']['raw_comments']:,} 条原始评论出发，经过文本清洗、规范化去重与 fastText 语言识别，形成 {english_total:,} 条英语评论的正式语料。随后使用 BERTopic 识别讨论主题，以三分类 Transformer 模型测量正面、中性和负面态度，并以 VADER 作为稳健性参照；最后通过词形还原后的滑动窗口共现网络揭示核心概念之间的关联。结果显示，英语评论的讨论同时涵盖作品的成长叙事、动画品质、产业比较、跨文化身份与平台互动；总体态度以正面评价为主，但中性说明与负面质疑亦占有相当比重。本文不把所选视频的最新评论视为全体 YouTube 用户意见，而将其作为跨文化数字接受的一项可复现、探索性样本。")
    add_body(document, "关键词：哪吒之魔童闹海；YouTube 评论；英语受众；BERTopic；情感分析；语义网络；计算传播")

    document.add_heading("一、研究问题与研究设计", level=1)
    add_body(document, "参考计算传播研究将大规模评论文本转化为可比较的内容、态度与关系结构的做法，本文提出四个相互关联的问题：RQ1，英语评论最集中讨论哪些主题；RQ2，评论的正面、中性与负面情感如何分布；RQ3，不同主题是否呈现不同的情感构成；RQ4，哪些内容词处于语义共现网络的中心。这样的设计避免仅凭词频或少数评论形成印象式判断，并允许从主题、情感和概念连接三个层面交叉解释受众接受。")
    add_body(document, "样本来自已筛选的 30 条《哪吒之魔童闹海》相关 YouTube 视频下的公开顶层评论。采集阶段未保留评论者显示名、头像、频道链接等可识别字段；分析阶段只使用匿名评论 ID、时间、互动量和正文。数据来源因此适合观察“所选视频评论区中的英语话语”，但不等同于 YouTube 全量历史评论，也不能推断全部英语受众的稳定意见。")

    document.add_heading("二、数据处理与方法", level=1)
    add_body(document, "数据处理采用平衡清洗原则：删除空白文本、少于 5 个字符的文本、纯链接、纯符号/数字和规范化重复评论；随后以 fastText 语言识别保留英语置信度不低于 0.80 的内容。非英语内容不进入主题、情感或网络分析，其匿名 ID、语言识别结果与排除原因仅记录在审计表中。处理漏斗保留了每一步数量，因而从英语语料与排除记录可以回溯至原始评论总数。")
    add_figure(document, "fig_01_cleaning_funnel.png", "图1 英语语料构建漏斗", "图1表明，原始评论经文本质量控制和语言识别后形成可追溯的英语分析语料；后续所有统计均以该语料为基础。")
    add_body(document, "主题分析使用 all-MiniLM-L6-v2 生成英文句向量，以 UMAP 降维、HDBSCAN 聚类和 BERTopic c-TF-IDF 表征形成主题。主题命名不是模型预设分类，而是结合主题规模、关键词和匿名代表性短引文进行的研究者解释。主题和情感关键词中均去除了英语停用词、代词、助动词、缩略形式，以及 ne、zha、movie、film、video、watch 等标题或平台泛词，使可视化更接近实质性评价与讨论内容。")
    add_body(document, "情感部分采用 cardiffnlp/twitter-roberta-base-sentiment-latest 的正面—中性—负面三分类输出；VADER 只用于稳健性对照，不替代主模型。网络部分保留名词、专有名词、形容词与动词的词形还原结果，在五词滑动窗口内统计共现，绘制高权重概念网络并报告加权度和介数中心性。月度结果仅作描述性展示，不用以作因果或长期演变推断。")

    document.add_heading("三、主题发现：英语受众在谈论什么（RQ1）", level=1)
    add_body(document, f"BERTopic 在 {english_total:,} 条英语评论中识别出 {len(visible_topics)} 个非离群主题；另有 {topic_meta['topic_counts'].get('-1', topic_meta['topic_counts'].get(-1, 0)):,} 条评论被归入离群/混合组，不被强行解释为单一主题。图2显示，规模最大的主题围绕观影价值与成长叙事，其次为动画品质、影评互动与观影推荐；中国电影与好莱坞的比较也形成一个清晰而规模可观的讨论簇。这说明英语评论区并非只给出简短好恶，而会把动画审美、文化位置与个人经验并置。")
    add_figure(document, "fig_02_topic_sizes.png", "图2 英语评论主要 BERTopic 主题的规模（图内仅用英文 Times New Roman 标签）", "图2显示，观影价值与成长叙事、动画品质与影评互动是最具规模的两类讨论，产业比较主题居于其后。")
    add_figure(document, "fig_03_topic_umap.png", "图3 英语评论主题空间的二维投影", "图3中相邻点表示语义表达较为接近；不同主题仍存在交叠，说明实际评论常同时涉及多个讨论面向。")
    topic_table = visible_topics[["topic_id", "comment_count", "中文主题", "keyword_label"]].rename(columns={"topic_id": "主题编号", "comment_count": "评论数", "中文主题": "解释性主题名称", "keyword_label": "模型关键词标签"})
    add_table(document, topic_table, ["主题编号", "评论数", "解释性主题名称", "模型关键词标签"])
    add_caption(document, "表1 英语评论主题概览")
    for _, theme in visible_topics.iterrows():
        add_body(document, theme_sentence(int(theme["topic_id"]), int(theme["comment_count"]), float(theme["comment_count"]) / visible_total))
        quote = quotes[quotes["topic_id"].eq(theme["topic_id"])].head(1)
        if not quote.empty:
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Cm(0.74)
            run = paragraph.add_run(f"匿名短引文：\"{quote.iloc[0]['quote']}\"")
            set_run_font(run, 9.5)

    document.add_heading("四、情感结构及其主题差异（RQ2—RQ3）", level=1)
    add_body(document, f"Transformer 分类显示，正面评论为 {int(sentiment_count.get('positive', 0)):,} 条（{sentiment_map.get('positive', 0):.1%}），中性评论为 {int(sentiment_count.get('neutral', 0)):,} 条（{sentiment_map.get('neutral', 0):.1%}），负面评论为 {int(sentiment_count.get('negative', 0)):,} 条（{sentiment_map.get('negative', 0):.1%}）。正面态度居首，说明该样本整体更倾向认可作品；但中性与负面合计超过半数，表明评论区并不是单向“叫好”，而是同时存在信息补充、比较说明、审美分歧和批评。Transformer 与 VADER 的逐条标签一致率为 {sentiment_meta['vader_agreement']:.1%}，其水平提示不同方法对网络语体、反讽和语境的处理并不完全一致，因而本文将情感结果理解为分布性指标，而非对单条评论的绝对判定。")
    add_figure(document, "fig_04_sentiment_distribution.png", "图4 英语评论的三分类情感分布", "图4表明正面评价占比最高，但中性说明和负面批评合计仍占较大部分，评论区呈现多元态度结构。")
    add_figure(document, "fig_05_monthly_sentiment.png", "图5 月度情感占比（仅作描述性展示）", "图5仅用于呈现样本在不同月份的构成差异；由于各月样本量和视频来源不同，不据此作趋势或因果推断。")
    add_body(document, f"在剔除代词、功能词、影片标题和平台泛词后，正面语料中较突出的内容词包括：{positive_terms}；负面语料中较突出的内容词包括：{negative_terms}。词云因此不再由 and、the、this、you、it 等语法词主导，而用于辅助展示两类态度的实质性评价焦点。词云是频率可视化，不应被视为词义或因果关系的证明。")
    add_figure(document, "fig_06_positive_wordcloud.png", "图6 正面评论的内容词云（已剔除代词、功能词、片名泛词与国家/地区词）", "图6突出动画、故事、优秀和惊艳等内容词，反映正面评价主要围绕制作、叙事与总体观影体验展开。")
    add_figure(document, "fig_06_negative_wordcloud.png", "图7 负面评论的内容词云（已剔除代词、功能词、片名泛词与国家/地区词）", "图7突出丑、坏、宣传、剧情等内容词，显示负面表达更多指向审美判断、叙事质疑或价值争议。")
    if not crosstab.empty:
        add_body(document, "为避免将总体情感分布误投射到每一个讨论焦点，本文进一步按主题计算三类情感的行内占比。图8展示了主题内部的情感构成，而非主题之间的绝对评论量差异。它可用于比较哪些话题更常伴随赞赏、解释性中立表达或批评；但主题类别来自无监督聚类，且各主题样本量不同，因此该图提供的是关联性描述，而非显著性检验或因果结论。")
        add_figure(document, "fig_09_topic_sentiment_heatmap.png", "图8 BERTopic 主题内部的情感构成（百分比）", "图8显示，不同主题并不共享同一种态度结构；动画品质主题更偏正面，而产业比较主题的负面比例相对更高。")

    document.add_heading("五、语义网络：核心概念及其连接（RQ4）", level=1)
    add_body(document, f"在去除代词、停用词、标题泛词和低频内容词后，语义网络保留 {network_meta['nodes']} 个节点、{network_meta['edges']} 条高权重共现边，并识别出 {network_meta['communities']} 个概念社群。加权度较高的概念包括 {top_network_terms}。其中，China/chinese 连接文化身份、产业比较与跨国传播的讨论；animation、story、character 等词则把制作质量、叙事与角色评价串联起来。网络中的边表示评论中词语在局部语境内共同出现的倾向，不表示概念之间具有因果方向。")
    add_figure(document, "fig_07_semantic_network.png", "图9 英语评论内容词的加权语义共现网络", "图9中节点越大、连线越密集，表示相应概念在评论语境中的共现联系越强；该关系是描述性的而非因果性的。")
    add_figure(document, "fig_08_network_centrality.png", "图10 语义网络中加权度最高的核心概念", "图10进一步量化了核心概念的连接强度，说明动画、故事等作品层面的词汇在评论语义结构中具有较高的组织作用。")
    centrality = network.head(15)[["term", "document_frequency", "weighted_degree", "betweenness"]].copy()
    centrality["betweenness"] = centrality["betweenness"].map(lambda value: f"{value:.3f}")
    centrality = centrality.rename(columns={"term": "概念", "document_frequency": "文档频次", "weighted_degree": "加权度", "betweenness": "介数中心性"})
    add_table(document, centrality, ["概念", "文档频次", "加权度", "介数中心性"], limit=15)
    add_caption(document, "表2 语义网络的核心概念")

    document.add_heading("六、讨论", level=1)
    add_body(document, "结合主题、情感与网络结果可以看到，英语受众的接受并非简单的“喜欢/不喜欢”二分。第一，成长、价值和身份隐喻构成最大主题，说明动画的情感与伦理叙事是跨文化理解的重要入口。第二，动画品质和剧情角色的讨论，与中国/好莱坞比较、语言发音和东亚身份话语并存，说明评论者一方面以普遍化的影视评价标准观看作品，另一方面也持续意识到其中国动画与跨文化传播的位置。第三，创作者互动主题提示平台语境不可忽略：评论区的对象既是电影，也包括发布评论视频的创作者。")
    add_body(document, "这种三层互证的解释方式与参考研究的计算传播逻辑相一致：主题模型回答“谈什么”，情感分类回答“以何种态度谈”，网络分析回答“概念如何联结”。三者的结果共同表明，英语评论中的作品接受同时具有审美评价、文化比较、身份协商与平台社交四种面向。特别是中国/chinese 在网络中的中心位置，并不能直接证明某种国家形象效应，却说明文化和产业定位是该样本中组织讨论的重要节点。")

    document.add_heading("七、局限与结论", level=1)
    add_body(document, "本研究存在四项主要局限。其一，样本来自经检索筛选的 30 条视频及其可获得的最新公开顶层评论，不能代表 YouTube 全量历史评论或所有英语受众。其二，fastText 语言识别阈值和文本清洗会排除部分混合语、极短评论或非标准拼写。其三，BERTopic 的主题数、聚类边界与自动情感模型均会受参数和语境影响；VADER 对照的一致率也表明模型间存在差异。其四，词云和语义共现网络都是描述性工具，不应被解释为因果机制或受众人口属性。后续研究可扩大视频样本、引入人工编码与跨平台语料，并对讽刺、文化典故和主题命名开展人工复核。")
    add_body(document, "总体而言，本文建立了一条可复现的英语 YouTube 评论分析流程，并在原始数据保持只读的前提下，给出了《哪吒之魔童闹海》跨文化数字接受的初步图景。结果显示，英语讨论以正面评价为主，但并不止于作品赞赏：成长叙事、动画与叙事品质、产业位置、文化身份和平台互动共同塑造了评论区的话语结构。")

    document.add_heading("附录：可复现性信息", level=1)
    versions = {}
    for package in ("pandas", "bertopic", "transformers", "sentence-transformers", "networkx", "spacy", "python-docx"):
        try:
            versions[package] = importlib.metadata.version(package)
        except importlib.metadata.PackageNotFoundError:
            versions[package] = "not found"
    reproducibility = {
        "generated_at": utc_now(),
        "raw_hashes": {"comments.csv": sha256(RAW_COMMENTS), "videos.csv": sha256(RAW_VIDEOS)},
        "preparation": prep, "topic_metadata": topic_meta, "sentiment_metadata": sentiment_meta,
        "network_metadata": network_meta, "package_versions": versions,
    }
    write_json(REPORTS_DIR / "reproducibility_metadata.json", reproducibility)
    add_body(document, "原始输入文件哈希、清洗漏斗、模型信息、随机种子、依赖版本和运行元数据已写入 reports/reproducibility_metadata.json。原始采集目录未被本分析流程修改。")
    # The original report may be open in Word; keep it untouched and write the
    # revised, expanded report to a distinct filename.
    output = REPORTS_DIR / "哪吒2英语YouTube评论计算分析报告_修订版.docx"
    document.save(output)
    print(f"Word report written to {output}")


if __name__ == "__main__":
    main()
