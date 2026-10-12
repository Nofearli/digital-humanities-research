"""Generate a Chinese academic-style Word report for Chinese comments."""

from __future__ import annotations

import importlib.metadata

import pandas as pd
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from pipeline_common import DATA_DIR, FIGURES_DIR, RAW_COMMENTS, RAW_VIDEOS, REPORTS_DIR, RESULTS_DIR, ensure_directories, read_json, sha256, utc_now, write_json


def set_run_font(run, size: float = 10.5, bold: bool = False) -> None:
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(size)
    run.bold = bold


def add_body(document: Document, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.first_line_indent = Cm(.74)
    paragraph.paragraph_format.line_spacing = 1.45
    paragraph.paragraph_format.space_after = Pt(5)
    set_run_font(paragraph.add_run(text), 10.5)


def add_caption(document: Document, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(3)
    set_run_font(paragraph.add_run(text), 10)


def add_figure(document: Document, filename: str, caption: str, explanation: str, width: float = 15.5) -> None:
    path = FIGURES_DIR / filename
    if path.exists():
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.add_run().add_picture(str(path), width=Cm(width))
        add_caption(document, caption)
        add_body(document, explanation)


def set_cell(cell, value: object) -> None:
    cell.text = str(value)
    for paragraph in cell.paragraphs:
        for run in paragraph.runs:
            set_run_font(run, 8.5)


def add_table(document: Document, frame: pd.DataFrame, columns: list[str], limit: int | None = None) -> None:
    frame = frame.loc[:, columns].head(limit) if limit else frame.loc[:, columns]
    table = document.add_table(rows=1, cols=len(columns))
    table.style = "Table Grid"
    for cell, name in zip(table.rows[0].cells, columns):
        set_cell(cell, name)
    for _, row in frame.iterrows():
        cells = table.add_row().cells
        for cell, name in zip(cells, columns):
            set_cell(cell, row[name])


def configure(document: Document) -> None:
    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)
    for name in ("Heading 1", "Heading 2"):
        style = document.styles[name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    for section in document.sections:
        section.top_margin = Cm(2.54)
        section.bottom_margin = Cm(2.54)
        section.left_margin = Cm(3)
        section.right_margin = Cm(3)


def main() -> None:
    ensure_directories()
    prep = read_json(DATA_DIR / "preparation_summary.json")
    topics = pd.read_csv(RESULTS_DIR / "topic_summary.csv", encoding="utf-8-sig")
    quotes = pd.read_csv(RESULTS_DIR / "topic_representative_quotes.csv", encoding="utf-8-sig")
    sentiment = pd.read_csv(RESULTS_DIR / "sentiment_summary.csv", encoding="utf-8-sig")
    keywords = pd.read_csv(RESULTS_DIR / "sentiment_keywords.csv", encoding="utf-8-sig")
    network = pd.read_csv(RESULTS_DIR / "network_nodes.csv", encoding="utf-8-sig")
    topic_meta = read_json(RESULTS_DIR / "topic_metadata.json")
    sentiment_meta = read_json(RESULTS_DIR / "sentiment_metadata.json")
    network_meta = read_json(RESULTS_DIR / "network_metadata.json")
    crosstab_path = RESULTS_DIR / "topic_sentiment_crosstab.csv"
    crosstab = pd.read_csv(crosstab_path, encoding="utf-8-sig") if crosstab_path.exists() else pd.DataFrame()
    chinese_total = prep["funnel"]["高置信度中文"]
    visible = topics[~topics["is_outlier"].astype(str).str.lower().eq("true")].sort_values("comment_count", ascending=False)
    sentiment_map = dict(zip(sentiment["sentiment"], sentiment["share"]))
    sentiment_count = dict(zip(sentiment["sentiment"], sentiment["comment_count"]))
    positive_terms = "、".join(keywords[keywords["sentiment"].eq("positive")].head(8)["term"].tolist())
    negative_terms = "、".join(keywords[keywords["sentiment"].eq("negative")].head(8)["term"].tolist())
    central_terms = "、".join(network.head(8)["term"].tolist())

    document = Document()
    configure(document)
    title = document.add_heading("《哪吒之魔童闹海》中文 YouTube 评论的计算文本分析", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in title.runs:
        set_run_font(run, 16, True)
    subtitle = document.add_paragraph("——主题、情感与语义网络的探索性研究报告")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in subtitle.runs:
        set_run_font(run, 11)

    document.add_heading("摘要", level=1)
    add_body(document, f"本研究以 YouTube 平台《哪吒之魔童闹海》相关视频下的公开顶层评论为材料，考察中文评论者的讨论主题、情感态度与核心概念结构。研究从 {prep['funnel']['原始评论']:,} 条原始评论出发，经文本质量控制、去重和 fastText 语言识别，保留 {chinese_total:,} 条语言置信度不低于 0.80 的中文评论。研究以 BERTopic 识别讨论主题，以中文原生 Transformer 的正向概率构建正面—中性—负面三分类，并以 SnowNLP 作稳健性对照；最后构建内容词共现网络。结果呈现中文评论中的多元话题与态度结构。本文将样本界定为所选视频评论区的探索性观察，不将其外推为全体中文受众意见。")
    add_body(document, "关键词：哪吒之魔童闹海；中文评论；YouTube；主题模型；情感分析；语义网络")

    document.add_heading("一、研究问题与数据来源", level=1)
    add_body(document, "研究围绕四个问题展开：RQ1，中文评论集中讨论哪些主题；RQ2，正面、中性和负面情感如何分布；RQ3，不同主题的情感构成是否存在差异；RQ4，哪些概念处于评论语义网络的中心。数据来自 30 条经筛选的相关视频下可获取的公开顶层评论。采集过程不保留评论者显示名、头像或频道链接；分析中只使用匿名评论 ID、时间、互动量和正文。")

    document.add_heading("二、数据处理与方法", level=1)
    add_body(document, "处理阶段删除空白文本、少于五字符的文本、纯链接、纯符号/数字及规范化重复评论。随后用 fastText 仅保留语言代码为 zh、置信度不低于 0.80 的内容。未进入正式语料的条目只在审计表记录匿名 ID、语言、置信度和排除原因，不保留正文。主题分析采用 BGE 中文句向量、jieba 分词和 BERTopic；为避免该语料初始密度聚类产生超大混合簇，使用固定八类 KMeans 形成语义簇，再以主题内平均 TF-IDF 提取关键词。网络分析保留名词、动词和形容词，并在五词窗口中统计共现。")
    add_figure(document, "fig_01_cleaning_funnel.png", "图1 中文语料构建漏斗", "图1呈现从原始评论到高置信度中文语料的处理过程，后续主题、情感与网络统计均以最终语料为基础。")
    add_body(document, "情感主模型为中文原生二分类 Transformer。正向概率不低于 0.60 标记为正面，不高于 0.40 标记为负面，位于中间区间的低置信度文本标记为中性；这一中性定义是概率阈值策略，而非独立训练的中性类别。SnowNLP 以相同阈值转化为三类，用于稳健性参照。")

    document.add_heading("三、主题发现：中文评论在谈论什么", level=1)
    outlier_count = topic_meta["topic_counts"].get("-1", topic_meta["topic_counts"].get(-1, 0))
    add_body(document, f"BERTopic 在 {chinese_total:,} 条中文评论中识别出 {len(visible)} 个非离群主题，另有 {outlier_count:,} 条归为离群或混合评论。主题名称由模型关键词、主题规模和匿名代表性评论综合解释，不是预先给定的类别。")
    add_figure(document, "fig_02_topic_sizes.png", "图2 中文评论主要 BERTopic 主题的规模", "图2比较各讨论簇的规模，能够识别评论者更常围绕哪些议题展开表达。")
    add_figure(document, "fig_03_topic_umap.png", "图3 中文评论主题空间二维投影", "图3中相邻点代表语义表达更接近；主题间的交叠说明单条评论可能同时关联多个议题。")
    table = visible[["topic_id", "comment_count", "interpreted_label", "keyword_label"]].rename(columns={"topic_id": "主题编号", "comment_count": "评论数", "interpreted_label": "主题标签", "keyword_label": "关键词"})
    add_table(document, table, ["主题编号", "评论数", "主题标签", "关键词"])
    add_caption(document, "表1 中文评论主题概览")
    for _, topic in visible.iterrows():
        add_body(document, f"主题 {int(topic['topic_id'])} 包含 {int(topic['comment_count']):,} 条评论，关键词为“{topic['keyword_label']}”。该标签反映这一讨论簇的高权重表达，具体解释应与匿名评论样例结合阅读。")
        quote = quotes[quotes["topic_id"].eq(topic["topic_id"])].head(1)
        if not quote.empty:
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Cm(.74)
            set_run_font(paragraph.add_run(f"匿名短引文：\"{quote.iloc[0]['quote']}\""), 9.5)

    document.add_heading("四、情感结构与主题差异", level=1)
    add_body(document, f"主模型将 {int(sentiment_count.get('positive', 0)):,} 条评论（{sentiment_map.get('positive', 0):.1%}）判为正面，将 {int(sentiment_count.get('neutral', 0)):,} 条评论（{sentiment_map.get('neutral', 0):.1%}）判为中性，将 {int(sentiment_count.get('negative', 0)):,} 条评论（{sentiment_map.get('negative', 0):.1%}）判为负面。与 SnowNLP 的三类标签一致率为 {sentiment_meta['snownlp_agreement']:.1%}。该对照用于提示结果的稳健边界，正文以 Transformer 概率阈值分类为主。")
    add_figure(document, "fig_04_sentiment_distribution.png", "图4 中文评论的三分类情感分布", "图4概括整体情感构成；其含义是样本中的模型分类比例，而不是对所有中文观众态度的总体估计。")
    add_figure(document, "fig_05_monthly_sentiment.png", "图5 月度情感占比（描述性展示）", "图5用于描述不同月份的样本构成。各月评论量与视频来源不同，因此不作长期趋势或因果解释。")
    add_body(document, f"去除功能词、代词、片名与平台泛词后，正面评论的高权重内容词包括：{positive_terms}；负面评论的高权重内容词包括：{negative_terms}。词云用于辅助展示评价焦点，不单独构成主题或因果结论。")
    add_figure(document, "fig_06_positive_wordcloud.png", "图6 正面评论内容词云", "图6突出正面语料中的高频内容词，辅助识别受认可的制作、叙事或情感体验。")
    add_figure(document, "fig_06_negative_wordcloud.png", "图7 负面评论内容词云", "图7突出负面语料中的高频内容词，辅助识别批评、质疑或争议的表述焦点。")
    if not crosstab.empty:
        add_figure(document, "fig_09_topic_sentiment_heatmap.png", "图8 不同 BERTopic 主题的情感构成（百分比）", "图8比较主题内部的三类情感比例，显示不同讨论焦点可能伴随不同的态度结构；该结果为描述性关联。")

    document.add_heading("五、语义网络：核心概念与连接", level=1)
    add_body(document, f"语义网络保留 {network_meta['nodes']} 个节点和 {network_meta['edges']} 条高权重共现边，识别出 {network_meta['communities']} 个概念社群。加权度较高的概念包括：{central_terms}。节点的加权度反映其与其他内容词在局部语境中的共现强度，边只表示关联而非因果关系。")
    add_figure(document, "fig_07_semantic_network.png", "图9 中文评论内容词的语义共现网络", "图9中节点越大、连线越密集，表示相关概念在评论文本中具有更强的共现联系。")
    add_figure(document, "fig_08_network_centrality.png", "图10 语义网络中加权度最高的核心概念", "图10将网络结构转化为可比较的中心性排名，帮助识别组织中文评论话语的关键概念。")
    centrality = network.head(15)[["term", "document_frequency", "weighted_degree", "betweenness"]].copy()
    centrality["betweenness"] = centrality["betweenness"].map(lambda value: f"{value:.3f}")
    centrality = centrality.rename(columns={"term": "概念", "document_frequency": "文档频次", "weighted_degree": "加权度", "betweenness": "介数中心性"})
    add_table(document, centrality, ["概念", "文档频次", "加权度", "介数中心性"], 15)
    add_caption(document, "表2 中文评论语义网络的核心概念")

    document.add_heading("六、讨论、局限与结论", level=1)
    add_body(document, "本文以内容、情感和结构三个层面分析中文评论。主题模型回答评论者谈论什么，情感模型描述态度构成，语义网络呈现概念之间的连接。三种方法互为补充，避免只依赖词频或少数评论形成印象式结论。对主题标签、词云和网络中心性的解释均应回到评论语境，而不能将其视为因果机制。")
    add_body(document, "研究的局限包括：样本来自所选视频的可获取公开顶层评论，不能代表 YouTube 全量中文评论；语言识别和清洗会排除部分极短或混合语评论；主题聚类、概率阈值和 SnowNLP 均可能产生误判；二分类模型的中性类别来自置信度区间，须谨慎解释。后续可扩展视频样本、加入人工编码与跨平台语料，检验主题和情感结论的稳定性。")
    add_body(document, "总体而言，本研究建立了独立、可复现的中文评论分析流程，并在原始数据保持只读的前提下，为《哪吒之魔童闹海》的中文数字接受提供主题、情感和语义结构的初步描述。")

    document.add_heading("附录：可复现性信息", level=1)
    versions = {}
    for package in ("pandas", "bertopic", "transformers", "sentence-transformers", "jieba", "snownlp", "networkx", "python-docx"):
        try:
            versions[package] = importlib.metadata.version(package)
        except importlib.metadata.PackageNotFoundError:
            versions[package] = "not found"
    write_json(REPORTS_DIR / "reproducibility_metadata.json", {"generated_at": utc_now(), "raw_hashes": {"comments.csv": sha256(RAW_COMMENTS), "videos.csv": sha256(RAW_VIDEOS)}, "preparation": prep, "topic_metadata": topic_meta, "sentiment_metadata": sentiment_meta, "network_metadata": network_meta, "package_versions": versions})
    add_body(document, "原始输入哈希、清洗漏斗、模型、阈值、依赖版本和运行元数据已写入 reports/reproducibility_metadata.json。原始采集目录及英文分析目录均未被本流程修改。")
    output = REPORTS_DIR / "哪吒2中文YouTube评论计算分析报告.docx"
    document.save(output)
    print(f"Word report written to {output}")


if __name__ == "__main__":
    main()
